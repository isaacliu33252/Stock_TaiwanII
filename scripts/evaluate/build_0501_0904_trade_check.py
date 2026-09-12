#!/usr/bin/env python3
"""Build a trade-by-trade review workbook for 0501_0904.xlsx."""

from __future__ import annotations

import argparse
import json
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

import duckdb
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_INPUT = PROJECT_ROOT / "0501_0904.xlsx"
DEFAULT_OUTPUT = PROJECT_ROOT / "0501_0904_check.xlsx"
DEFAULT_DB = PROJECT_ROOT / "FinRL/data/stock_data.db"
DEFAULT_LIVE_SIGNAL = PROJECT_ROOT / "report/group_a_plus/latest/live_signal.json"

NAME_TO_TICKER = {
    "元大台灣50": "0050.TW",
    "元大台灣50正2": "00631L.TW",
    "元大台灣50反1": "00632R.TW",
    "元大美債20年": "00679B.TWO",
    "元大AAA至A公司債": "00751B.TWO",
    "元大高股息": "0056.TW",
    "國泰永續高股息": "00878.TW",
    "元大台灣高息低波": "00713.TW",
    "元大S&P500": "00646.TW",
    "玉山金": "2884.TW",
}

GROUP_A_PLUS_TICKERS = {"0050.TW", "00631L.TW", "00632R.TW", "00679B.TWO"}


def _resolve(raw: str | Path) -> Path:
    path = Path(raw)
    return path if path.is_absolute() else PROJECT_ROOT / path


def _num(value: Any) -> float | None:
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _load_live_signal(path: Path) -> dict[str, Any]:
    payload = json.loads(path.read_text(encoding="utf-8-sig"))
    data = payload.get("data")
    return data if isinstance(data, dict) else payload


def _latest_closes(db_path: Path, tickers: list[str], as_of: str) -> dict[str, dict[str, Any]]:
    con = duckdb.connect(str(db_path), read_only=True)
    try:
        rows = con.execute(
            """
            SELECT ticker, dt, close
            FROM (
                SELECT
                    ticker,
                    dt,
                    close,
                    ROW_NUMBER() OVER (PARTITION BY ticker ORDER BY dt DESC) AS rn
                FROM ohlcv
                WHERE ticker IN ({placeholders}) AND dt <= ?
            )
            WHERE rn = 1
            """.format(placeholders=",".join(["?"] * len(tickers))),
            [*tickers, as_of],
        ).fetchall()
    finally:
        con.close()
    return {str(ticker): {"dt": str(dt), "close": float(close)} for ticker, dt, close in rows}


def _read_trades(path: Path) -> list[dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    trades: list[dict[str, Any]] = []
    for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if not any(value is not None for value in row):
            continue
        trade_date = row[0]
        if isinstance(trade_date, datetime):
            trade_date_s = trade_date.date().isoformat()
        elif isinstance(trade_date, date):
            trade_date_s = trade_date.isoformat()
        else:
            trade_date_s = str(trade_date)
        name = str(row[2]).strip()
        qty = _num(row[3]) or 0.0
        price = _num(row[4]) or 0.0
        broker_amount = _num(row[5])
        fee = _num(row[6]) or 0.0
        net_cash_flow = _num(row[10])
        gross = qty * price
        side_raw = str(row[1]).strip()
        side = "buy" if side_raw in {"現股買進", "定期定額"} else "sell" if side_raw == "現股賣出" else "unknown"
        trades.append(
            {
                "row_no": idx,
                "date": trade_date_s,
                "trade_type": side_raw,
                "side": side,
                "name": name,
                "ticker": NAME_TO_TICKER.get(name, ""),
                "quantity": qty,
                "price": price,
                "broker_amount": broker_amount,
                "computed_gross": gross,
                "gross_minus_broker_amount": None if broker_amount is None else gross - broker_amount,
                "fee": fee,
                "net_cash_flow": net_cash_flow,
            }
        )
    return trades


def _alignment(side: str, ticker: str, target_weight: float) -> str:
    if ticker not in GROUP_A_PLUS_TICKERS:
        return "off_universe"
    if side == "buy" and target_weight > 0:
        return "aligned_add"
    if side == "buy" and target_weight <= 0:
        return "against_target_buy_zero_weight"
    if side == "sell" and target_weight > 0:
        return "against_target_sell_positive_weight"
    if side == "sell" and target_weight <= 0:
        return "aligned_reduce_zero_weight"
    return "unknown"


def _recommendation(alignment: str, ticker: str, side: str, actual_vs_hold: float | None) -> tuple[str, str]:
    if alignment == "aligned_add":
        return ("keep_but_stage", "符合 latest strategy 正權重；建議分批、受 turnover/POV 限制，不追價。")
    if alignment == "aligned_reduce_zero_weight":
        return ("keep_reduce", "符合 latest strategy 零權重；賣出可降低非目標曝險。")
    if alignment == "against_target_sell_positive_weight":
        return ("avoid_sell_core", "latest strategy 仍需此正權重；除非風控或現金需求，避免賣核心部位。")
    if alignment == "against_target_buy_zero_weight":
        return ("avoid_buy_zero_weight", "latest strategy 目標為 0；只允許 shadow/短期對沖理由，不宜常態買進。")
    if alignment == "off_universe":
        return ("redirect_to_core_or_cash", "不屬 GroupA+ latest universe；新增資金優先轉向 0050/00631L 或保留 cash。")
    if side == "buy" and actual_vs_hold is not None and actual_vs_hold < 0:
        return ("tighten_entry_filter", "買後至 9/4 低於成本；需要等待策略確認或分批降低進場風險。")
    return ("manual_review", "需人工確認交易目的。")


def _analyze(
    trades: list[dict[str, Any]],
    *,
    target_weights: dict[str, float],
    closes: dict[str, dict[str, Any]],
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for trade in trades:
        ticker = trade["ticker"]
        close_info = closes.get(ticker, {})
        close = _num(close_info.get("close"))
        target_weight = float(target_weights.get(ticker, 0.0))
        gross = float(trade["computed_gross"])
        fee = float(trade["fee"])
        side = str(trade["side"])
        qty = float(trade["quantity"])
        actual_vs_hold = None
        interpretation = ""
        if close is not None and side == "buy":
            actual_vs_hold = qty * close - gross - fee
            interpretation = "buy_pnl_to_asof"
        elif close is not None and side == "sell":
            actual_vs_hold = gross - qty * close - fee
            interpretation = "sell_cash_vs_holding_to_asof"
        alignment = _alignment(side, ticker, target_weight)
        action, recommendation = _recommendation(alignment, ticker, side, actual_vs_hold)
        rows.append(
            {
                **trade,
                "latest_close_date": close_info.get("dt"),
                "latest_close": close,
                "latest_target_weight": target_weight,
                "in_group_a_plus": ticker in GROUP_A_PLUS_TICKERS,
                "alignment": alignment,
                "actual_vs_hold_to_asof": actual_vs_hold,
                "actual_vs_hold_pct": None if actual_vs_hold is None or gross == 0 else actual_vs_hold / gross,
                "metric_interpretation": interpretation,
                "suggested_action": action,
                "improvement_note": recommendation,
            }
        )
    return rows


def _summary(rows: list[dict[str, Any]], live: dict[str, Any]) -> list[list[Any]]:
    total_buy = sum(r["computed_gross"] + r["fee"] for r in rows if r["side"] == "buy")
    total_sell = sum(r["computed_gross"] - r["fee"] for r in rows if r["side"] == "sell")
    off_universe_buy = sum(r["computed_gross"] + r["fee"] for r in rows if r["side"] == "buy" and not r["in_group_a_plus"])
    zero_weight_buy = sum(
        r["computed_gross"] + r["fee"]
        for r in rows
        if r["side"] == "buy" and r["in_group_a_plus"] and r["latest_target_weight"] <= 0
    )
    core_sell = sum(
        r["computed_gross"] - r["fee"]
        for r in rows
        if r["side"] == "sell" and r["in_group_a_plus"] and r["latest_target_weight"] > 0
    )
    scored = [r for r in rows if r["actual_vs_hold_to_asof"] is not None]
    buy_pnl = sum(r["actual_vs_hold_to_asof"] for r in scored if r["side"] == "buy")
    sell_vs_hold = sum(r["actual_vs_hold_to_asof"] for r in scored if r["side"] == "sell")
    return [
        ["input_period", "2026-05-08 to 2026-09-04"],
        ["latest_strategy_requested_as_of", live.get("requested_as_of_date")],
        ["latest_strategy_actual_data_date", live.get("actual_data_date")],
        ["latest_strategy_id", live.get("strategy_id")],
        ["trade_count", len(rows)],
        ["buy_notional_plus_fee", round(total_buy, 2)],
        ["sell_notional_minus_fee", round(total_sell, 2)],
        ["off_universe_buy_notional", round(off_universe_buy, 2)],
        ["zero_weight_group_a_plus_buy_notional", round(zero_weight_buy, 2)],
        ["positive_weight_group_a_plus_sell_notional", round(core_sell, 2)],
        ["buy_pnl_to_2026_09_04_close", round(buy_pnl, 2)],
        ["sell_cash_vs_holding_to_2026_09_04_close", round(sell_vs_hold, 2)],
        ["main_improvement_1", "停止 off-universe 定期定額/加碼；新增資金先對齊 0050/00631L/cash。"],
        ["main_improvement_2", "00632R 與 00679B latest target 為 0；只在正式風控 gate 開啟時使用。"],
        ["main_improvement_3", "不要在 latest strategy 仍正權重時賣出 0050/00631L，除非有現金或風控理由。"],
        ["main_improvement_4", "所有交易先做 pre-trade check：目標權重差、現金 30% floor、turnover、稅費、POV。"],
    ]


def _by_ticker(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[str, dict[str, Any]] = defaultdict(lambda: defaultdict(float))
    names: dict[str, str] = {}
    for r in rows:
        key = r["ticker"] or r["name"]
        names[key] = r["name"]
        g = grouped[key]
        if r["side"] == "buy":
            g["buy_count"] += 1
            g["buy_quantity"] += r["quantity"]
            g["buy_notional"] += r["computed_gross"] + r["fee"]
        elif r["side"] == "sell":
            g["sell_count"] += 1
            g["sell_quantity"] += r["quantity"]
            g["sell_notional"] += r["computed_gross"] - r["fee"]
        if r["actual_vs_hold_to_asof"] is not None:
            g["actual_vs_hold_to_asof"] += r["actual_vs_hold_to_asof"]
        g["off_universe_count"] += 0 if r["in_group_a_plus"] else 1
        g["against_count"] += 1 if str(r["alignment"]).startswith("against") or r["alignment"] == "off_universe" else 0
    output = []
    for ticker, values in sorted(grouped.items()):
        output.append({"ticker": ticker, "name": names[ticker], **{k: round(v, 4) for k, v in values.items()}})
    return output


def _alignment_summary(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[str, dict[str, Any]] = defaultdict(lambda: defaultdict(float))
    notes = {
        "aligned_add": "買進 latest strategy 正權重標的，原則可保留但要分批與控成本。",
        "aligned_reduce_zero_weight": "賣出 latest strategy 零權重標的，方向正確。",
        "against_target_sell_positive_weight": "賣出 latest strategy 正權重核心，通常降低未來收益。",
        "against_target_buy_zero_weight": "買進 latest strategy 零權重標的，需停止常態化。",
        "off_universe": "不在 GroupA+ latest universe，會分散資金與決策焦點。",
    }
    for row in rows:
        key = row["alignment"]
        group = grouped[key]
        group["trade_count"] += 1
        if row["side"] == "buy":
            group["buy_notional"] += row["computed_gross"] + row["fee"]
        elif row["side"] == "sell":
            group["sell_notional"] += row["computed_gross"] - row["fee"]
        if row["actual_vs_hold_to_asof"] is not None:
            group["actual_vs_hold_to_asof"] += row["actual_vs_hold_to_asof"]
    output = []
    for key, values in sorted(grouped.items()):
        output.append(
            {
                "alignment": key,
                "trade_count": int(values["trade_count"]),
                "buy_notional": round(values["buy_notional"], 2),
                "sell_notional": round(values["sell_notional"], 2),
                "actual_vs_hold_to_asof": round(values["actual_vs_hold_to_asof"], 2),
                "improvement": notes.get(key, "manual review"),
            }
        )
    return output


def _data_quality(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    output: list[dict[str, Any]] = []
    for row in rows:
        issues: list[str] = []
        diff = row.get("gross_minus_broker_amount")
        if diff is not None and abs(float(diff)) > 2.0:
            issues.append("broker_amount_differs_from_quantity_times_price")
        if not row.get("ticker"):
            issues.append("ticker_mapping_missing")
        if row.get("latest_close") is None:
            issues.append("latest_close_missing")
        if issues:
            output.append(
                {
                    "row_no": row["row_no"],
                    "date": row["date"],
                    "name": row["name"],
                    "ticker": row["ticker"],
                    "broker_amount": row["broker_amount"],
                    "computed_gross": row["computed_gross"],
                    "gross_minus_broker_amount": row["gross_minus_broker_amount"],
                    "issues": "; ".join(issues),
                }
            )
    return output


def _append_table(ws, rows: list[dict[str, Any]] | list[list[Any]]) -> None:
    if not rows:
        return
    if isinstance(rows[0], dict):
        headers = list(rows[0].keys())
        ws.append(headers)
        for row in rows:
            ws.append([row.get(h) for h in headers])
    else:
        for row in rows:
            ws.append(row)


def _format_workbook(wb: Workbook) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        for col in range(1, ws.max_column + 1):
            letter = get_column_letter(col)
            max_len = 10
            for cell in ws[letter]:
                if cell.value is not None:
                    max_len = max(max_len, min(60, len(str(cell.value)) + 2))
            ws.column_dimensions[letter].width = max_len
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)


def build_workbook(*, input_path: Path, output_path: Path, db_path: Path, live_signal_path: Path) -> None:
    live = _load_live_signal(live_signal_path)
    target_weights = {str(k): float(v) for k, v in (live.get("target_weights") or {}).items()}
    trades = _read_trades(input_path)
    tickers = sorted({ticker for ticker in NAME_TO_TICKER.values()})
    closes = _latest_closes(db_path, tickers, str(live.get("actual_data_date") or "2026-09-04"))
    rows = _analyze(trades, target_weights=target_weights, closes=closes)

    wb = Workbook()
    ws = wb.active
    ws.title = "summary"
    _append_table(ws, _summary(rows, live))

    ws = wb.create_sheet("trade_check")
    _append_table(ws, rows)

    ws = wb.create_sheet("by_ticker")
    _append_table(ws, _by_ticker(rows))

    ws = wb.create_sheet("alignment_summary")
    _append_table(ws, _alignment_summary(rows))

    ws = wb.create_sheet("data_quality")
    _append_table(ws, _data_quality(rows))

    ws = wb.create_sheet("latest_strategy")
    strategy_rows = []
    for ticker, weight in target_weights.items():
        strategy_rows.append(
            {
                "ticker": ticker,
                "target_weight": weight,
                "target_value": (live.get("target_values") or {}).get(ticker),
                "reference_target_shares_before_cost": (live.get("reference_target_shares_before_cost") or {}).get(ticker),
                "latest_close_date": closes.get(ticker, {}).get("dt"),
                "latest_close": closes.get(ticker, {}).get("close"),
            }
        )
    _append_table(ws, strategy_rows)

    ws = wb.create_sheet("rules")
    _append_table(
        ws,
        [
            ["rule", "meaning"],
            ["aligned_add", "買進 latest strategy 正權重標的。"],
            ["aligned_reduce_zero_weight", "賣出 latest strategy 零權重標的。"],
            ["against_target_sell_positive_weight", "賣出 latest strategy 仍需持有的正權重標的。"],
            ["against_target_buy_zero_weight", "買進 latest strategy 目標為 0 的 GroupA+ 標的。"],
            ["off_universe", "交易不在 GroupA+ latest universe 中，會分散資金與決策焦點。"],
            ["actual_vs_hold_to_asof", "買進為持有到 2026-09-04 的估算損益；賣出為賣出現金相對續抱到 2026-09-04 的差額。"],
        ],
    )

    _format_workbook(wb)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", default=str(DEFAULT_INPUT))
    parser.add_argument("--output", default=str(DEFAULT_OUTPUT))
    parser.add_argument("--db", default=str(DEFAULT_DB))
    parser.add_argument("--live-signal", default=str(DEFAULT_LIVE_SIGNAL))
    args = parser.parse_args()
    build_workbook(
        input_path=_resolve(args.input),
        output_path=_resolve(args.output),
        db_path=_resolve(args.db),
        live_signal_path=_resolve(args.live_signal),
    )
    print(f"wrote {_resolve(args.output)}")


if __name__ == "__main__":
    main()
